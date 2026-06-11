"""Excel import/export for entire Atlas worlds."""
from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Tuple

from services.atlas import canvas as atlas_canvas
from services.atlas import csv_io as atlas_csv
from services.atlas import documents as atlas_documents
from services.atlas import entities as atlas_entities
from services.atlas import relationships as atlas_relationships
from services.atlas.worlds import get_world, create_world, refresh_world_stats

META_SHEET = "_Atlas"
META_VERSION = 2
EXPORT_ROW_LIMIT = 10000

_INVALID_SHEET_CHARS = re.compile(r'[\[\]:*?/\\]')


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
        from openpyxl import Workbook, load_workbook  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Atlas Excel import/export. "
            "Install it with: pip install openpyxl"
        ) from exc
    from openpyxl import Workbook, load_workbook
    return Workbook, load_workbook


def _utcnow_iso() -> str:
    return datetime.now(timezone.utc).replace(tzinfo=None).isoformat()


def _slugify_name(name: str) -> str:
    s = (name or "").strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    return s or "untitled"


def _sanitize_sheet_name(name: str, used: set) -> str:
    base = _INVALID_SHEET_CHARS.sub("", (name or "Sheet").strip())[:31] or "Sheet"
    candidate = base
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = f"{base[: max(1, 31 - len(suffix))]}{suffix}"
        n += 1
    used.add(candidate)
    return candidate


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _parse_cell_value(raw: Any) -> Any:
    if raw is None or raw == "":
        return None
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, int) and not isinstance(raw, bool):
        return raw
    if isinstance(raw, float):
        if raw.is_integer():
            return int(raw)
        return raw
    s = str(raw).strip()
    if not s:
        return None
    if s.lower() in ("true", "false"):
        return s.lower() == "true"
    if s.startswith("{") or s.startswith("["):
        try:
            return json.loads(s)
        except json.JSONDecodeError:
            pass
    if re.fullmatch(r"-?\d+", s):
        try:
            return int(s)
        except ValueError:
            pass
    if re.fullmatch(r"-?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            pass
    return s


def _similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, _slugify_name(a), _slugify_name(b)).ratio()


def _read_meta(wb) -> Optional[Dict[str, Any]]:
    if META_SHEET not in wb.sheetnames:
        return None
    ws = wb[META_SHEET]
    raw = ws.cell(1, 1).value
    if not raw:
        return None
    try:
        return json.loads(str(raw))
    except json.JSONDecodeError:
        return None


def _sheet_rows(ws) -> Tuple[List[str], List[Dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    headers = [h for h in headers if h]
    if not headers:
        return [], []
    out: List[Dict[str, Any]] = []
    for row in rows[1:]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        item: Dict[str, Any] = {}
        for i, key in enumerate(headers):
            if i >= len(row):
                break
            val = _parse_cell_value(row[i])
            if val is not None:
                item[key] = val
        out.append(item)
    return headers, out


def export_world(owner: Optional[str], world_id: str) -> Tuple[bytes, Dict[str, Any]]:
    Workbook, _ = _require_openpyxl()
    world = get_world(owner, world_id)
    entities = atlas_entities.list_entities(owner, world_id)
    layout = atlas_canvas.get_layout(owner, world_id)
    rels = atlas_relationships.list_relationships(owner, world_id)

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set = set()

    entity_meta: List[Dict[str, Any]] = []
    total_docs = 0

    for ent in entities:
        sheet_name = _sanitize_sheet_name(ent["name"], used_names)
        ws = wb.create_sheet(sheet_name)
        result = atlas_documents.find(owner, world_id, ent["id"], limit=EXPORT_ROW_LIMIT, offset=0)
        docs = result["documents"]
        total_docs += len(docs)
        all_keys: set = set()
        for d in docs:
            all_keys.update(k for k in d if not k.startswith("_") or k == "_atlas_row_id")
        slugs = ["_atlas_row_id"] + sorted(k for k in all_keys if k != "_atlas_row_id")
        ws.append(slugs)
        for d in docs:
            row = []
            for s in slugs:
                v = d.get(s, d.get("_id") if s == "_atlas_row_id" else "")
                row.append(_cell_str(v))
            ws.append(row)
        entity_meta.append({
            "id": ent["id"],
            "name": ent["name"],
            "sheet_name": sheet_name,
        })

    meta = {
        "version": META_VERSION,
        "world_id": world_id,
        "world_name": world["name"],
        "description": world.get("description") or "",
        "exported_at": _utcnow_iso(),
        "entities": entity_meta,
        "relationships": rels,
        "canvas": layout,
    }
    meta_ws = wb.create_sheet(META_SHEET)
    meta_ws.cell(1, 1, json.dumps(meta, ensure_ascii=False))
    meta_ws.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    stats = {
        "entity_count": len(entities),
        "document_count": total_docs,
        "relationship_count": len(rels),
        "world_name": world["name"],
    }
    return buf.getvalue(), stats


def _match_entity(
    sheet_name: str,
    entities: List[Dict[str, Any]],
    meta: Optional[Dict[str, Any]],
) -> Dict[str, Any]:
    meta_entities = {e.get("sheet_name", ""): e for e in (meta or {}).get("entities") or []}
    meta_by_name = {e.get("name", "").lower(): e for e in (meta or {}).get("entities") or []}

    if sheet_name in meta_entities:
        me = meta_entities[sheet_name]
        for ent in entities:
            if ent["id"] == me.get("id") or ent["name"].lower() == (me.get("name") or "").lower():
                return {
                    "suggested_action": "import_to_existing",
                    "suggested_entity_id": ent["id"],
                    "suggested_entity_name": ent["name"],
                    "match_confidence": "exact",
                    "alternatives": [],
                }

    for ent in entities:
        if ent["name"].lower() == sheet_name.lower():
            return {
                "suggested_action": "import_to_existing",
                "suggested_entity_id": ent["id"],
                "suggested_entity_name": ent["name"],
                "match_confidence": "exact",
                "alternatives": [],
            }

    slug_sheet = _slugify_name(sheet_name)
    for ent in entities:
        if _slugify_name(ent["name"]) == slug_sheet:
            return {
                "suggested_action": "import_to_existing",
                "suggested_entity_id": ent["id"],
                "suggested_entity_name": ent["name"],
                "match_confidence": "exact",
                "alternatives": [],
            }

    scored: List[Tuple[float, Dict[str, Any]]] = []
    for ent in entities:
        ratio = _similarity(sheet_name, ent["name"])
        if ratio >= 0.85:
            scored.append((ratio, ent))
    scored.sort(key=lambda x: -x[0])
    if scored:
        best = scored[0][1]
        alts = [{"entity_id": e["id"], "entity_name": e["name"], "score": round(s, 3)} for s, e in scored[1:4]]
        return {
            "suggested_action": "import_to_existing",
            "suggested_entity_id": best["id"],
            "suggested_entity_name": best["name"],
            "match_confidence": "fuzzy",
            "alternatives": alts,
        }

    if sheet_name.lower() in meta_by_name:
        me = meta_by_name[sheet_name.lower()]
        return {
            "suggested_action": "create_new",
            "suggested_entity_id": None,
            "suggested_entity_name": me.get("name") or sheet_name,
            "match_confidence": "meta",
            "alternatives": [],
        }

    return {
        "suggested_action": "create_new",
        "suggested_entity_id": None,
        "suggested_entity_name": sheet_name,
        "match_confidence": "new",
        "alternatives": [{"entity_id": e["id"], "entity_name": e["name"]} for e in entities[:5]],
    }


def analyze_import(
    owner: Optional[str],
    file_bytes: bytes,
    world_id: Optional[str] = None,
) -> Dict[str, Any]:
    _, load_workbook = _require_openpyxl()
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    meta = _read_meta(wb)
    entities: List[Dict[str, Any]] = []
    if world_id:
        entities = atlas_entities.list_entities(owner, world_id)

    sheets: List[Dict[str, Any]] = []
    for name in wb.sheetnames:
        if name == META_SHEET:
            continue
        ws = wb[name]
        headers, rows = _sheet_rows(ws)
        match = _match_entity(name, entities, meta)
        sheets.append({
            "sheet_name": name,
            "row_count": len(rows),
            "columns": headers,
            **match,
        })

    world_hint = None
    if meta:
        world_hint = {
            "world_name": meta.get("world_name"),
            "description": meta.get("description"),
            "entity_count": len(meta.get("entities") or []),
            "relationship_count": len(meta.get("relationships") or []),
        }

    return {
        "world_id": world_id,
        "world_hint": world_hint,
        "sheets": sheets,
        "has_meta": meta is not None,
    }


def _rows_to_csv(headers: List[str], rows: List[Dict[str, Any]]) -> str:
    import csv
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=headers, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: _cell_str(row.get(k, "")) for k in headers})
    return buf.getvalue()


def _import_sheet_rows(
    owner: Optional[str],
    world_id: str,
    entity_id: str,
    headers: List[str],
    rows: List[Dict[str, Any]],
    mode: str,
) -> Dict[str, Any]:
    if not rows:
        return {"rows_total": 0, "rows_inserted": 0, "rows_updated": 0, "rows_skipped": 0, "rows_failed": 0, "errors": []}
    csv_text = _rows_to_csv(headers, rows)
    return atlas_csv.import_rows(owner, world_id, entity_id, csv_text, mode=mode)


def import_world(
    owner: Optional[str],
    file_bytes: bytes,
    mappings: List[Dict[str, Any]],
    *,
    world_id: Optional[str] = None,
    world_name: Optional[str] = None,
    world_description: str = "",
    default_mode: str = "append",
    restore_meta: bool = True,
) -> Dict[str, Any]:
    _, load_workbook = _require_openpyxl()
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    meta = _read_meta(wb) if restore_meta else None

    id_map: Dict[str, str] = {}
    result: Dict[str, Any] = {
        "world_id": world_id,
        "entities_created": 0,
        "entities_updated": 0,
        "rows_inserted": 0,
        "rows_updated": 0,
        "rows_failed": 0,
        "relationships_restored": 0,
        "sheet_results": [],
        "errors": [],
    }

    if not world_id:
        wname = world_name or (meta or {}).get("world_name") or "Imported World"
        wdesc = world_description or (meta or {}).get("description") or ""
        world = create_world(owner, wname, wdesc)
        world_id = world["id"]
        result["world_id"] = world_id

    mapping_by_sheet = {m["sheet_name"]: m for m in mappings}

    for sheet_name in wb.sheetnames:
        if sheet_name == META_SHEET:
            continue
        m = mapping_by_sheet.get(sheet_name)
        if not m or m.get("action") == "skip":
            continue
        ws = wb[sheet_name]
        headers, rows = _sheet_rows(ws)
        mode = m.get("mode") or default_mode
        action = m.get("action") or "create_new"
        entity_id = m.get("entity_id")
        entity_name = m.get("entity_name") or sheet_name

        try:
            if action == "import_to_existing" and entity_id:
                ent = atlas_entities.get_entity(owner, world_id, entity_id)
                old_id = ent["id"]
            else:
                ent = atlas_entities.create_entity(owner, world_id, entity_name)
                old_id = None
                result["entities_created"] += 1
                entity_id = ent["id"]

            stats = _import_sheet_rows(owner, world_id, entity_id, headers, rows, mode)
            result["rows_inserted"] += stats.get("rows_inserted", 0)
            result["rows_updated"] += stats.get("rows_updated", 0)
            result["rows_failed"] += stats.get("rows_failed", 0)
            if stats.get("errors"):
                result["errors"].extend(stats["errors"])
            result["sheet_results"].append({"sheet_name": sheet_name, "entity_id": entity_id, **stats})

            meta_ent = next(
                (e for e in (meta or {}).get("entities") or [] if e.get("sheet_name") == sheet_name),
                None,
            )
            if meta_ent and meta_ent.get("id"):
                id_map[meta_ent["id"]] = entity_id
            elif old_id:
                id_map[old_id] = entity_id
        except Exception as e:
            result["errors"].append({"sheet": sheet_name, "message": str(e)})

    if restore_meta and meta:
        for rel in meta.get("relationships") or []:
            from_id = id_map.get(rel.get("from_entity_id"), rel.get("from_entity_id"))
            to_id = id_map.get(rel.get("to_entity_id"), rel.get("to_entity_id"))
            try:
                atlas_entities.get_entity(owner, world_id, from_id)
                atlas_entities.get_entity(owner, world_id, to_id)
                atlas_relationships.create_relationship(
                    owner, world_id,
                    from_id, to_id,
                    rel.get("rel_type", "one_to_many"),
                    rel.get("from_field", ""),
                    rel.get("to_field", ""),
                    label=rel.get("label", ""),
                    from_anchor=rel.get("from_anchor", "right"),
                    to_anchor=rel.get("to_anchor", "left"),
                )
                result["relationships_restored"] += 1
            except Exception as e:
                result["errors"].append({"relationship": rel.get("label") or rel.get("id"), "message": str(e)})

        canvas = meta.get("canvas") or {}
        nodes = []
        for n in canvas.get("nodes") or []:
            old_eid = n.get("entity_id")
            new_eid = id_map.get(old_eid, old_eid)
            try:
                atlas_entities.get_entity(owner, world_id, new_eid)
                nodes.append({**n, "entity_id": new_eid})
            except Exception:
                continue
        clusters = canvas.get("clusters") or []
        if nodes or clusters:
            atlas_canvas.save_layout(owner, world_id, nodes, clusters=clusters)

    refresh_world_stats(owner, world_id)
    return result
